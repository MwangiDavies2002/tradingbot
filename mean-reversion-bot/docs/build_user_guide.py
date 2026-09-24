"""Build the operator guide using the Codex document runtime.

Preset: compact_reference_guide; header: editorial_cover (compact opening).
Named overrides: Title 26pt/navy; Subtitle 12pt/gray; Code 9pt Consolas;
running furniture 9pt gray. No tables or decorative borders.
"""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = Path(__file__).parent / 'Mean_Reversion_Bot_User_Guide.docx'
doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(1)
s.header_distance = s.footer_distance = Inches(.492)
for name, size, color, before, after in [
    ('Normal',11,'202A35',0,6), ('Title',26,'1F4D78',0,10),
    ('Subtitle',12,'596675',0,12), ('Heading 1',16,'2E74B5',18,10),
    ('Heading 2',13,'2E74B5',14,7), ('Heading 3',12,'1F4D78',10,5),
    ('List Bullet',11,'202A35',0,4), ('List Number',11,'202A35',0,4),
    ('Header',9,'596675',0,0), ('Footer',9,'596675',0,0),
]:
    st = doc.styles[name]
    st.font.name = 'Calibri'; st.font.size = Pt(size)
    st.font.color.rgb = RGBColor.from_string(color)
    pf = st.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = 1.25; pf.widow_control = True
    if name.startswith('Heading'): pf.keep_with_next = True
# Explicit real Word numbering geometry, including hanging indents and tab stops.
for lvl in doc.part.numbering_part.element.findall('.//' + qn('w:lvl')):
    pp = lvl.find(qn('w:pPr'))
    if pp is None: pp = OxmlElement('w:pPr'); lvl.append(pp)
    for old in list(pp): pp.remove(old)
    ind = OxmlElement('w:ind'); ind.set(qn('w:left'),'540'); ind.set(qn('w:hanging'),'271'); pp.append(ind)
    tabs = OxmlElement('w:tabs'); tab = OxmlElement('w:tab')
    tab.set(qn('w:val'),'num'); tab.set(qn('w:pos'),'540'); tabs.append(tab); pp.append(tabs)
    spacing = OxmlElement('w:spacing')
    for k,v in [('before','0'),('after','80'),('line','300'),('lineRule','auto')]: spacing.set(qn('w:'+k),v)
    pp.append(spacing)
s.header.paragraphs[0].text = 'MEAN-REVERSION BOT  |  MT5 DEMO USER GUIDE'
foot = s.footer.paragraphs[0]; foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
foot.add_run('24 September 2026  |  Page ')
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); foot._p.append(fld)
def p(t, style=None): return doc.add_paragraph(t, style)
def h(t): doc.add_heading(t,1)
def sub(t): doc.add_heading(t,2)
def bullet(t): p(t,'List Bullet')
def step(t): p(t,'List Number')
def page(title): doc.add_page_break(); h(title)

t=p('Mean-Reversion Bot','Title'); t.alignment=WD_ALIGN_PARAGRAPH.CENTER
t=p('Practical user guide • Connected MT5 workflow','Subtitle'); t.alignment=WD_ALIGN_PARAGRAPH.CENTER
p('For the local Windows application, based on the code present on 24 September 2026. Your MT5 connection is assumed to be set up; this guide does not certify its current status.')
h('Start here: your first session')
p('The MT5 runner supports demo accounts and Volatility 75 (1s) Index only. Its lab symbol is 1HZ75V. It rejects real accounts. Connecting the terminal does not start entries.')
step('Open http://localhost:3000/backtest on the Windows PC running MT5 and the backend. Sign in with your app access key if prompted; this is separate from your MT5 login.')
step('In Strategy Lab, change Platform from TradingView to MT5 / Python. Check the connected account and server. If disconnected, click Connect MT5 demo.')
step('Click Load saved selection to restore the configuration saved for the runner. Review the timeframe, Active Strategies and Minimum score.')
step('Review Risk per trade (%) and Daily equity loss limit (%). Click Save lab selection after any changes. Confirm the saved summary matches your intended settings.')
step('Click Start demo. Check that the panel says Running, then monitor its status message and Last poll. A qualifying completed candle is needed before an entry can occur.')
step('Use MT5 Demo & Journal to review positions and decisions. To pause new entries, click Stop entries and confirm Stopped.')
sub('Know what stopping does')
p('Stop entries does not close positions or cancel existing pending orders. Existing positions keep their broker stop loss and take profit. An order already being submitted may finish. Manage any remaining exposure in MT5 and verify it there.')

page('Configure the strategy you intend to run')
p('Use Strategy Lab for editing. The separate MT5 Demo & Journal page provides monitoring and connect/start/stop controls, but does not expose the lab selection editor.')
sub('Signals and timeframe')
bullet('Timeframe: M1, M5, M15, M30, H1 or H4. The runner evaluates closed candles, not each price tick as a new signal.')
bullet('Active Strategies: toggle the signal components you want included. Available controls include Z-Score, RSI, Bollinger Bands, VWAP, Stochastic, liquidity/structure, volume, Hurst and other filters.')
bullet('Minimum score: an integer from 1 to 20. This is a weighted point threshold, not a count of agreeing indicators. Direction checks and regime filters can still block an entry.')
bullet('MT5 execution is fixed to Volatility 75 (1s) Index. Selecting another asset for research does not change the MT5 runner’s instrument.')
sub('Risk controls')
p('Risk per trade defaults to 0.5%, with a maximum of 1%. The planned budget uses the lower of account balance and equity. The worker calculates broker lots in account currency, rounds down to the lot step and skips the entry if the minimum lot exceeds the budget.')
p('The entry includes an ATR-based stop loss and a take-profit distance of twice the stop distance, subject to broker checks and price rounding. Spread, costs, slippage or gaps can make actual loss exceed the planned stop risk. These defaults are software settings, not validated profit parameters.')
p('Daily equity loss limit defaults to 3%, with a maximum of 5%. Its baseline is account equity at the first worker poll of the UTC day. The account-wide equity check can be affected by other trading. Once triggered, new entries remain blocked for that UTC day, including after a restart. It does not close open positions.')
sub('Save, then start')
p('Stop the runner before changing its saved strategy. Edit in the lab, click Save lab selection, verify the saved summary, then Start demo. Browser-restored controls and the backend’s saved selection can differ; Load saved selection reconciles them. Starting after a reload uses the backend’s saved configuration.')

page('Monitor, pause and review a session')
sub('What a normal session looks like')
p('Keep Windows awake, the desktop terminal logged in, and the backend running. The browser may be closed without stopping the worker. Refreshing or signing out of the UI is not a stop command.')
p('The worker polls about every two seconds; the panel refreshes about every five seconds. It needs at least 100 current closed candles and handles each candle once. Running means monitoring is active; it does not mean a trade is due or guaranteed.')
bullet('Check account/server, Running or Stopped, the status message, and Last poll. A stale page is not evidence that trading has stopped.')
bullet('Review Balance and Equity, then the open positions table: ticket, side, lots, entry, SL/TP and floating P&L. Confirm positions and protective orders in MT5 Toolbox → Trade.')
bullet('Any position or pending order on this symbol, including a manual trade, blocks a new bot entry. The worker does not stack additional exposure on the symbol.')
sub('Use the journal')
p('Expand Trade journal on MT5 Demo & Journal. Signal records include the reason, weighted score and strategy snapshot. Order requests and broker responses show what was attempted; deal records show executions and exits.')
p('Click Download complete journal CSV to save the full journal, beyond the latest events displayed. It contains timestamp_utc, account, kind and details_json. Net deal P&L includes profit, commission, swap and fee. Review all related deals for a position when assessing its full result.')
p('Bot-position deal history, including manual exits, is synchronized while running, when viewing the stopped connected journal, and after reconnecting. Check MT5 History if the app has not yet caught up. Journal times use UTC; Nairobi is UTC+3.')
sub('Finish or recover')
p('Click Stop entries and wait for Stopped. Check remaining positions and pending orders in MT5, then export the journal. After a backend restart, reconnect, review the saved selection and start explicitly; entries do not automatically resume. After an uncertain order response, inspect the terminal and journal before restarting because the broker may already have executed it.')

page('Test ideas and understand the other pages')
sub('Run a historical test')
bullet('In Strategy Lab choose MT5 / Python to reveal the test controls. Select one asset, timeframe, signal toggles, Minimum score, test period and starting capital.')
bullet('Click Run Combined Test for the current data pipeline, or Import CSV / Excel to test your own candles. Uploading a file starts a test with the current controls.')
bullet('Review Total P&L, Win Rate, Trades, Max DD, Sharpe Ratio, the curves and available trade details. Compare results across different periods and retain losing runs as well as winning ones.')
p('Choose Historical data source explicitly: Deriv history or Connected MT5 history (V75 1s). MT5 history requires the local connected demo terminal and locks the instrument to 1HZ75V. A missing-history error does not fall back to Deriv. Import CSV / Excel uses your supplied file instead of either history source. New results and saved-run details identify the effective source; older runs without that record show Not recorded.')
p('CSV columns: timestamp, open, high, low, close, volume. Volume is optional but matters to volume-based signals. Use Unix seconds or explicit UTC timestamps such as 2026-09-01T00:00:00Z. Supply consistent OHLC values, the selected instrument/timeframe, and enough candles for indicator warm-up. Excel support depends on the installed reader dependencies.')
p('Backtests use a simulator rather than exact MT5 broker execution and lot sizing. The lab’s starting capital is simulated; it does not change your MT5 balance. Saving a research result does not start trading.')
sub('Keep the workflows separate')
bullet('TradingView: script/chart workflow. Features marked Python / MT5 only do not have an exact Pine port. A TradingView view does not start the local MT5 worker.')
bullet('Dashboard, Deriv History, Deriv Signals and Deriv Risk concern the separate Deriv workflow. Their data and controls are not proof of MT5 activity.')
bullet('Research validation records experiments and outcomes. Institutional lab runs offline analyses, including costs, risk, allocation, order lifecycle and auto-quoting simulations. These are not broker order controls.')
bullet('Instrument catalog stores supplied specifications and revisions for offline plans. Registering an instrument does not make it tradable through the V75-only MT5 runner.')
p('High-impact news only and per-indicator Pine translation choices are not saved into the MT5 demo strategy. Do not assume they constrain MT5 entries; use the saved MT5 summary as the configuration reference.')

page('Troubleshooting')
sub('Connected, but no trades')
p('Read the status message and recent signal reasons first. Check Running, a current Last poll, sufficient closed candles, the weighted threshold, regime/direction filters, and whether a position or pending order already exists. A strategy can run normally without finding entries. Do not treat the absence of trades as proof of a fault.')
sub('Start demo is unavailable or saving fails')
p('Select MT5 / Python in Strategy Lab, connect the demo terminal and save any changed selection. If already running, stop before saving. Verify your app role permits changes; a viewer cannot perform trading control actions. Read the displayed API error rather than repeatedly clicking Start.')
sub('Terminal disconnected or trading permission denied')
p('Check the intended demo login/server in the desktop terminal. Enable Algo Trading. Under Tools → Options → Expert Advisors, clear Disable automated trading via external Python API. Stop and reconnect after an account change or terminal error. Real-money accounts are rejected by this runner.')
sub('Symbol unavailable, stale candles or history missing')
p('Use the exact Volatility 75 (1s) Index symbol on the correct server. Open its chart in Market Watch and allow history to load; check the connection and PC clock. R_75 is a different instrument. For an MT5-history request, load more chart history or request fewer days. A backtest error should be resolved before relying on its output.')
sub('Minimum lot skipped, broker rejection or uncertain response')
p('A minimum-lot skip means the broker’s smallest allowed size exceeds the configured risk budget. An order-check rejection can reflect stop-distance or broker constraints. Review the exact message. An uncertain submission stops the worker without an automatic retry; check open positions, orders, History and journal records before taking further action.')
sub('Daily loss block or stop still finishing')
p('The daily block lasts until the next UTC day and survives restarts. Stop entries may wait for an in-flight terminal call. If the UI cannot confirm Stopped, inspect MT5 directly and investigate the backend logs; do not assume the request completed.')
sub('Local UI or backend unavailable')
p('MT5 controls require localhost on the Windows PC with the terminal. A hosted website cannot control this desktop integration. If localhost fails, use the startup steps on the next page and inspect logs. If a port is already occupied, verify it belongs to the intended app before launching another process.')

page('Restarting the app and keeping records')
sub('Normal restart on this PC')
p('If the app is already running, keep using that instance. Otherwise, open PowerShell in the mean-reversion-bot project folder (inside mean-reversion-bot_2) and run:')
code=p('powershell -ExecutionPolicy Bypass -File .\\start-local.ps1')
for r in code.runs: r.font.name='Consolas'; r.font.size=Pt(9)
p('Open http://localhost:3000/backtest. The launcher starts the API and frontend in the background and writes logs under logs/. It does not place orders. Check the existing service if it reports port 8000 or 3000 already in use.')
p('The local backend health endpoint is http://127.0.0.1:8000/health. A healthy API does not prove MT5 is connected or the worker is running; verify the MT5 panel separately.')
sub('If dependencies are missing')
p('Follow MT5_SETUP.md in the project. The launcher expects backend/.venv-mt5 with requirements-mt5.txt installed and frontend dependencies installed using npm ci. Use the MT5-capable environment, a single backend process and no auto-reload. No reinstall is needed for an already working setup.')
p('If terminal discovery selects the wrong installation or fails, set MT5_TERMINAL_PATH to the intended terminal64.exe before launching the backend. Keep MT5 credentials in the terminal; the app uses that logged-in session.')
sub('Files worth keeping')
bullet('backend/data/mt5_journal.sqlite3: saved MT5 strategy, signal/order/deal journal and persistent daily-entry guard state.')
bullet('backend/data/local.db: the separate local application database used by run_local.py.')
bullet('logs/backend.err.log and logs/backend.out.log: backend errors and output. Frontend logs are in the same logs folder.')
bullet('Downloaded journal CSVs and experiment exports: records for reviewing sessions and comparing strategy versions.')
p('Stop entries, verify remaining broker exposure, then stop the backend before copying its database files for backup. Treat the journal as sensitive account activity. Deleting it also removes safeguards and history; it is not a normal way to reset a session.')
sub('Guide scope and implementation references')
p('This guide describes the current local implementation, not a profitability assessment. It was checked against MT5Panel.tsx, Backtest.tsx, mt5_demo.py, the MT5 and backtest API routes, start-local.ps1 and MT5_SETUP.md. It does not claim a live terminal connection or a broker order was tested while preparing the document.')

doc.core_properties.title='Mean-Reversion Bot User Guide'
doc.core_properties.subject='Operating the local MT5 demo runner'
doc.core_properties.author='Mean-Reversion Bot Project'
OUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT)
# Structural checks complement, but do not replace, rendered visual QA.
check=Document(OUT)
assert len(check.paragraphs)>80
assert check.sections[0].page_width==Inches(8.5)
assert len(check.tables)==0
assert 'Stop entries' in '\n'.join(x.text for x in check.paragraphs)
print(f'Created {OUT}; paragraphs={len(check.paragraphs)}')
